"""Cobo 对客方案 deck 组件库（python-pptx）。

用法：先用 pptx skill 的 add_slide.py 从基准 deck 复制一页带标题的内容页，
重排、清理完成后，再用本模块在空白内容区画组件。依赖：pip install python-pptx
    import os, sys
    sys.path.insert(0, os.path.expanduser('~/.claude/skills/cobo-customer-solution-deck/scripts'))
    from deck_components import *
    p = Presentation('deck.pptx'); s = p.slides[3]
    clear_body(s); set_title(s, '平台你自己建，底層交給 Cobo')
    subtitle(s, '一句話副標')
    card_grid(s, [('卖点名', '一句客户收益'), ...], cols=3)
画布 10 x 5.625 英寸；正文区约 y 1.0–4.95；左边距 0.56；卡片间距 0.16。
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

BLUE  = RGBColor(0x00, 0x24, 0xA1)   # Cobo 主蓝：标题、卖点名、编号圆、强调条
DARK  = RGBColor(0x1A, 0x1D, 0x24)   # 正文
GREY  = RGBColor(0x59, 0x59, 0x59)   # 副标、脚注
TINT  = RGBColor(0xEE, 0xF1, 0xFB)   # 卡片底色
WATER = RGBColor(0xE8, 0xEE, 0xFB)   # 大面积浅底（如冰山水面）
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEAD  = RGBColor(0x9E, 0x9E, 0x9E)   # 表头（沿用原模板表格）
FIRST = RGBColor(0xF3, 0xF3, 0xF3)   # 表格首列
GREEN = RGBColor(0x2E, 0x7D, 0x32); AMBER = RGBColor(0xE0, 0x8E, 0x00); RED = RGBColor(0xC6, 0x28, 0x28)
# 红绿黄只用于风险分级、痛点后果，不做装饰色。

X0, GAP, BODY_W = 0.56, 0.16, 8.9

def _is_title(sh):
    return sh.has_text_frame and sh.top is not None and sh.top < Inches(0.4) and sh.width > Inches(5)

def set_title(s, text):
    """改标题文字，保留第一个 run 的格式；标题为空时新建 run，继承母版格式。"""
    ts = [sh for sh in s.shapes if _is_title(sh)]
    if not ts: raise ValueError('这一页找不到标题框（顶部 0.4 英寸内、宽于 5 英寸的文本框）；请从带标题的内容页复制')
    ps = ts[0].text_frame.paragraphs; rs = ps[0].runs
    if rs: rs[0].text = text
    else: ps[0].add_run().text = text
    for r in rs[1:]: r._r.getparent().remove(r._r)
    for pp in ps[1:]: pp._p.getparent().remove(pp._p)

def clear_body(s):
    """删掉除标题外的所有形状（母版里的页脚、logo 不受影响）。"""
    for sh in list(s.shapes):
        if not _is_title(sh): sh._element.getparent().remove(sh._element)

def text(s, x, y, w, h, paras, size=11, color=DARK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """paras: str | [run,...] | [[run,...], ...]；run = str 或 (str, {'size','bold','color'})"""
    if isinstance(paras, str): paras = [[paras]]
    elif paras and not isinstance(paras[0], list): paras = [paras]
    b = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = b.text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for k, runs in enumerate(paras):
        pa = tf.paragraphs[0] if k == 0 else tf.add_paragraph(); pa.alignment = align
        if k: pa.space_before = Pt(4)
        for r in runs:
            t, o = (r, {}) if isinstance(r, str) else r
            run = pa.add_run(); run.text = t; f = run.font
            f.size = Pt(o.get('size', size)); f.bold = o.get('bold', bold); f.color.rgb = o.get('color', color)
    return b

def subtitle(s, t):
    return text(s, X0, 0.92, BODY_W, 0.35, [(t, {'size': 14, 'color': GREY})])

def box(s, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.08, line=None):
    b = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.solid(); b.fill.fore_color.rgb = fill; b.shadow.inherit = False
    if line: b.line.color.rgb = line; b.line.width = Pt(1)
    else: b.line.fill.background()
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE: b.adjustments[0] = adj
    return b

def badge(s, x, y, d, label, fill=BLUE, fc=WHITE, size=12):
    """编号圆：1/2/3 或 Q。"""
    c = box(s, x, y, d, d, fill, MSO_SHAPE.OVAL); tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    r = tf.paragraphs[0].add_run(); r.text = str(label); tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = fc
    return c

def card(s, x, y, w, h, n, label, desc):
    """卖点卡：编号圆 + 卖点名（蓝粗 12.5）+ 一句收益（10.5）。"""
    box(s, x, y, w, h, TINT)
    badge(s, x + 0.18, y + 0.2, 0.36, n)
    text(s, x + 0.66, y + 0.2, w - 0.84, 0.36, [(label, {'bold': True, 'size': 12.5, 'color': BLUE})], anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.18, y + 0.66, w - 0.36, h - 0.78, [(desc, {'size': 10.5})])

def card_grid(s, items, cols=3, y=1.4, h=1.2, rowgap=0.12, width=BODY_W):
    """3 张一行；最后一行不满时自动拉宽（如 5 张 = 3 + 2）。最多 6 张。"""
    assert len(items) <= 6, '一页最多 6 个卖点'
    rows = [items[i:i + cols] for i in range(0, len(items), cols)]; n = 1
    for r, row in enumerate(rows):
        w = (width - GAP * (len(row) - 1)) / len(row)
        for c, (a, b) in enumerate(row):
            card(s, X0 + c * (w + GAP), y + r * (h + rowgap), w, h, n, a, b); n += 1
    return y + len(rows) * (h + rowgap)

def key_strip(s, y, head, points, h=0.78):
    """深蓝关键信息条，例如「為什麼不用通道？」+ 三个理由。每页最多一条。"""
    box(s, X0, y, BODY_W, h, BLUE)
    text(s, X0 + 0.22, y, 1.7, h, [(head, {'bold': True, 'size': 12.5, 'color': WHITE})], anchor=MSO_ANCHOR.MIDDLE)
    cw = (BODY_W - 2.1) / len(points)
    for k, t in enumerate(points):
        text(s, X0 + 2.0 + k * cw, y, cw - 0.15, h, [(t, {'size': 10.5, 'color': WHITE})], anchor=MSO_ANCHOR.MIDDLE)

def qa_grid(s, qa, y=1.2, h=1.62):
    """2x2 问答卡：客户原话提问 + 一句回答。"""
    w = (BODY_W - GAP) / 2
    for k, (q, a) in enumerate(qa):
        r, c = divmod(k, 2); x = X0 + c * (w + GAP); yy = y + r * (h + GAP)
        box(s, x, yy, w, h, TINT); badge(s, x + 0.2, yy + 0.22, 0.42, 'Q')
        text(s, x + 0.75, yy + 0.2, w - 0.95, 0.46, [(q, {'bold': True, 'size': 13, 'color': BLUE})], anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.75, yy + 0.8, w - 0.95, 0.7, [(a, {'size': 11.5})])

def tiers(s, rows, y=1.5, h=0.82):
    """风险分级条：[(级别, 动作, 说明, 颜色)]，颜色用 GREEN/AMBER/RED。"""
    for k, (lv, act, desc, col) in enumerate(rows):
        yy = y + k * (h + 0.16)
        box(s, X0, yy, BODY_W, h, TINT); box(s, X0, yy, 1.6, h, col)
        text(s, X0, yy, 1.6, h, [(lv, {'bold': True, 'size': 14, 'color': WHITE})], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(s, X0 + 1.89, yy, 2.2, h, [(act, {'bold': True, 'size': 14, 'color': BLUE})], anchor=MSO_ANCHOR.MIDDLE)
        text(s, X0 + 4.14, yy, 4.6, h, [(desc, {'size': 12})], anchor=MSO_ANCHOR.MIDDLE)

def steps(s, items, y=1.5, h=2.55):
    """横向流程：[(步骤名, 发生什么, 后果)]，后果用红字。适合痛点页。"""
    n = len(items); w = (BODY_W - GAP * (n - 1)) / n
    for k, (a, b, c) in enumerate(items):
        x = X0 + k * (w + GAP); box(s, x, y, w, h, TINT); badge(s, x + (w - 0.42) / 2, y + 0.2, 0.42, k + 1)
        text(s, x + 0.1, y + 0.72, w - 0.2, 0.35, [(a, {'bold': True, 'size': 13, 'color': BLUE})], align=PP_ALIGN.CENTER)
        text(s, x + 0.12, y + 1.2, w - 0.24, 0.5, [(b, {'size': 10.5})], align=PP_ALIGN.CENTER)
        if c: text(s, x + 0.12, y + 1.72, w - 0.24, 0.6, [('→ ' + c, {'size': 10.5, 'bold': True, 'color': RED})], align=PP_ALIGN.CENTER)
        if k < n - 1:
            a_ = s.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x + w + 0.01), Inches(y + 0.28), Inches(0.13), Inches(0.26))
            a_.fill.solid(); a_.fill.fore_color.rgb = BLUE; a_.line.fill.background()

def table(s, rows, colw, x=X0, y=1.05, rowh=0.34, size=10.5, center_from=1, hl=1):
    """对比表：首行灰底白字、首列浅灰粗体、细灰框线。✓/△/✗ 用文字，表下加图例。
    hl：多列矩阵（4 列以上）里要加粗变蓝的列，默认第 1 列（Cobo 列）；不是 Cobo 对比表时传 None。"""
    for o in [sh for sh in s.shapes if sh.has_table]: o._element.getparent().remove(o._element)
    nr, nc = len(rows), len(rows[0])
    gf = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(sum(colw)), Inches(rowh * nr)); tbl = gf.table
    tp = tbl._tbl.tblPr
    for a in ('firstRow', 'bandRow'): tp.set(a, '0')
    sid = tp.find(qn('a:tableStyleId'))
    if sid is not None: sid.text = '{5940675A-B579-460E-94D1-54222C63F5DA}'
    for k, w in enumerate(colw): tbl.columns[k].width = Inches(w)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(rowh)
        for j, val in enumerate(row):
            c = tbl.cell(i, j); c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.margin_left = c.margin_right = Inches(0.06); c.margin_top = c.margin_bottom = Inches(0.02)
            pa = c.text_frame.paragraphs[0]; c.text_frame.word_wrap = True
            pa.alignment = PP_ALIGN.CENTER if (center_from is not None and j >= center_from and nc > 3) else PP_ALIGN.LEFT
            r = pa.add_run(); r.text = val; f = r.font; f.size = Pt(size)
            fill = HEAD if i == 0 else FIRST if j == 0 else WHITE
            c.fill.solid(); c.fill.fore_color.rgb = fill
            f.bold = i == 0 or j == 0; f.color.rgb = WHITE if i == 0 else DARK
            if i > 0 and hl is not None and j == hl and nc > 3: f.color.rgb = BLUE; f.bold = True   # 多列矩阵里 Cobo 列加粗蓝
            tcPr = c._tc.get_or_add_tcPr()
            for tag in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
                ln = tcPr.makeelement(qn(tag), {'w': '9525'}); sf = ln.makeelement(qn('a:solidFill'), {})
                sf.append(sf.makeelement(qn('a:srgbClr'), {'val': 'CCCCCC'})); ln.append(sf); tcPr.append(ln)
            for e in [e for e in tcPr if e.tag in (qn('a:solidFill'), qn('a:noFill'))]: tcPr.remove(e); tcPr.append(e)
    return gf

def footnote(s, t, y=4.6):
    return text(s, X0, y, BODY_W, 0.3, [(t, {'size': 9.5, 'color': GREY})])
